import sqlite3
import re
import openpyxl
from pathlib import Path

DB_PATH = Path("/app/data/catalog.db") if Path("/app/data").exists() else Path(__file__).parent / "catalog.db"
DATA_DIR = Path(__file__).parent

# Mapeamento: nome do arquivo → segmento
ARQUIVOS_SEGMENTOS = {
    "+ VENDIDOS - BRINQUEDOS.xlsx": "BRINQUEDOS",
    "+ VENDIDOS - BRINQUEDOS EDUCATIVOS.xlsx": "BRINQUEDOS EDUCATIVOS",
    "+ VENDIDOS - CONVENIÊNCIA.xlsx": "CONVENIÊNCIA",
    "+ VENDIDOS - FARMÁCIA.xlsx": "FARMÁCIA",
    "+ VENDIDOS - PAPELARIA.xlsx": "PAPELARIA",
    "+ VENDIDOS - SUPERMERCADOS.xlsx": "SUPERMERCADOS",
}

# Mapeamento de aba → região
def _extrair_regiao(sheet_name: str) -> str:
    """Extrai a região do nome da aba. Ex: 'BRINQUEDOS - NORTE' → 'NORTE'"""
    parts = sheet_name.split(" - ", 1)
    return parts[-1].strip().upper() if len(parts) > 1 else sheet_name.strip().upper()


def _extrair_sku_nome(produto_str: str) -> tuple[str, str]:
    """
    Extrai SKU e nome do produto da string.
    Ex: 'SQ250 - SQUISHY - BICHINHOS DE APERTAR - 7898706186875'
    → sku='SQ250', nome='SQUISHY - BICHINHOS DE APERTAR'
    """
    if not produto_str:
        return "", ""
    produto_str = str(produto_str).strip()
    # SKU é tudo antes do primeiro ' - '
    parts = produto_str.split(" - ", 1)
    sku = parts[0].strip()
    resto = parts[1].strip() if len(parts) > 1 else ""
    # Remove código de barras do final (sequência de 13+ dígitos no final)
    nome = re.sub(r'\s*-?\s*\d{13,}$', '', resto).strip()
    # Remove ' - ' solto no final
    nome = nome.rstrip(' -').strip()
    return sku, nome


def build_catalog():
    """Lê as 6 planilhas de segmentos e grava em SQLite."""
    conn = sqlite3.connect(DB_PATH)
    conn.execute("DROP TABLE IF EXISTS produtos")
    conn.execute("DROP TABLE IF EXISTS segmentos_abas")

    conn.execute("CREATE TABLE segmentos_abas (nome TEXT)")
    conn.execute("""
        CREATE TABLE produtos (
            segmento TEXT,
            sku TEXT,
            produto TEXT,
            regiao TEXT,
            qtd_total INTEGER
        )
    """)
    conn.execute("CREATE INDEX IF NOT EXISTS idx_seg ON produtos(segmento)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_regiao ON produtos(regiao)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_sku ON produtos(sku)")

    segmentos_inseridos = []
    total_produtos = 0

    for arquivo, segmento in ARQUIVOS_SEGMENTOS.items():
        filepath = DATA_DIR / arquivo
        if not filepath.exists():
            print(f"[WARN] Arquivo nao encontrado: {filepath}")
            continue

        print(f"Processando: {arquivo} -> {segmento}")
        segmentos_inseridos.append(segmento)

        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        rows = []

        for sheet_name in wb.sheetnames:
            regiao = _extrair_regiao(sheet_name)
            ws = wb[sheet_name]

            for i, row in enumerate(ws.iter_rows(min_row=1, values_only=True)):
                # Pula cabeçalho
                if i == 0:
                    continue
                if not row or not row[0]:
                    continue

                produto_str = str(row[0]).strip()
                qtd = row[1] if len(row) > 1 and row[1] else 0

                try:
                    qtd = int(qtd)
                except (ValueError, TypeError):
                    qtd = 0

                sku, nome = _extrair_sku_nome(produto_str)
                if not sku:
                    continue

                rows.append((segmento, sku, nome, regiao, qtd))
                total_produtos += 1

                if len(rows) >= 1000:
                    conn.executemany("INSERT INTO produtos VALUES (?,?,?,?,?)", rows)
                    rows = []

        if rows:
            conn.executemany("INSERT INTO produtos VALUES (?,?,?,?,?)", rows)

        wb.close()

    # Insere segmentos
    conn.executemany("INSERT INTO segmentos_abas VALUES (?)", [(s,) for s in sorted(segmentos_inseridos)])

    conn.commit()
    conn.close()
    print(f"Catalogo construido: {DB_PATH} ({total_produtos} produtos de {len(segmentos_inseridos)} segmentos)")


def get_connection():
    return sqlite3.connect(DB_PATH)


def listar_segmentos() -> list[str]:
    """Retorna os segmentos em ordem alfabetica."""
    conn = get_connection()
    rows = conn.execute("SELECT nome FROM segmentos_abas ORDER BY nome").fetchall()
    conn.close()
    return [r[0] for r in rows if r[0]]


def listar_regioes(segmento: str | None = None) -> list[str]:
    conn = get_connection()
    if segmento:
        rows = conn.execute(
            "SELECT DISTINCT regiao FROM produtos WHERE segmento = ? ORDER BY regiao",
            (segmento,),
        ).fetchall()
    else:
        rows = conn.execute(
            "SELECT DISTINCT regiao FROM produtos ORDER BY regiao"
        ).fetchall()
    conn.close()
    return [r[0] for r in rows if r[0]]


def recomendar_produtos(
    segmento: str,
    regiao: str | None = None,
    uf: str | None = None,
    limite: int = 20,
) -> list[dict]:
    """Retorna os produtos mais vendidos para o segmento/regiao, ordenados por qtd_total."""
    conn = get_connection()
    query = """
        SELECT sku, produto, regiao,
               SUM(qtd_total) as total_vendas
        FROM produtos
        WHERE segmento = ?
    """
    params: list = [segmento]

    if regiao:
        query += " AND regiao = ?"
        params.append(regiao)

    query += """
        GROUP BY sku, produto
        ORDER BY total_vendas DESC
        LIMIT ?
    """
    params.append(limite)

    rows = conn.execute(query, params).fetchall()
    conn.close()

    return [
        {
            "sku": r[0],
            "produto": r[1],
            "regiao": r[2],
            "total_vendas": r[3],
        }
        for r in rows
    ]


if __name__ == "__main__":
    build_catalog()
