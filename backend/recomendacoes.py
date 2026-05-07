import json
import sqlite3
from datetime import datetime
from pathlib import Path
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

DB_PATH = Path("/app/data/recomendacoes.db") if Path("/app/data").exists() else Path(__file__).parent / "recomendacoes.db"


def _conn():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = _conn()
    conn.execute("""
        CREATE TABLE IF NOT EXISTS recomendacoes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            data TEXT NOT NULL,
            session_id TEXT,
            segmento TEXT,
            regiao TEXT,
            uf TEXT,
            porte TEXT,
            produtos_json TEXT,
            total REAL DEFAULT 0,
            qtd_produtos INTEGER DEFAULT 0
        )
    """)
    conn.commit()
    conn.close()


def salvar(session_id: str, segmento: str, regiao: str, uf: str, porte: str, produtos: list[dict], total: float):
    conn = _conn()
    conn.execute(
        """INSERT INTO recomendacoes (data, session_id, segmento, regiao, uf, porte, produtos_json, total, qtd_produtos)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            session_id,
            segmento or "",
            regiao or "",
            uf or "",
            porte or "",
            json.dumps(produtos, ensure_ascii=False),
            total,
            len(produtos),
        ),
    )
    conn.commit()
    conn.close()


def listar(limit: int = 100, offset: int = 0) -> list[dict]:
    conn = _conn()
    rows = conn.execute(
        "SELECT id, data, segmento, regiao, uf, porte, total, qtd_produtos FROM recomendacoes ORDER BY id DESC LIMIT ? OFFSET ?",
        (limit, offset),
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def obter(rec_id: int) -> dict | None:
    conn = _conn()
    row = conn.execute("SELECT * FROM recomendacoes WHERE id = ?", (rec_id,)).fetchone()
    conn.close()
    if not row:
        return None
    d = dict(row)
    d["produtos"] = json.loads(d.pop("produtos_json", "[]"))
    return d


def exportar_excel(rec_id: int) -> BytesIO | None:
    rec = obter(rec_id)
    if not rec:
        return None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Recomendacao"

    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="E63946", end_color="E63946", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    label_font = Font(bold=True, size=11)

    # Cabecalho do pedido
    info = [
        ("Recomendacao", f"#{rec['id']}"),
        ("Data", rec["data"]),
        ("Segmento", rec["segmento"]),
        ("Regiao", rec["regiao"]),
        ("UF", rec["uf"]),
        ("Porte", rec["porte"]),
        ("Total", f"R$ {rec['total']:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")),
    ]
    for i, (label, value) in enumerate(info, 1):
        ws.cell(row=i, column=1, value=label).font = label_font
        ws.cell(row=i, column=2, value=value)

    # Tabela de produtos
    start_row = len(info) + 2
    headers = ["#", "Produto", "Formato", "Preco Unit.", "Qtd", "Subtotal"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for i, p in enumerate(rec["produtos"], 1):
        row = start_row + i
        try:
            preco = float(p.get("preco", 0))
        except (ValueError, TypeError):
            preco = 0
        qty = int(p.get("quantidade", 1))
        subtotal = preco * qty

        ws.cell(row=row, column=1, value=i).border = thin_border
        ws.cell(row=row, column=2, value=p.get("nome", "")).border = thin_border
        ws.cell(row=row, column=3, value=p.get("formato", "")).border = thin_border
        c_preco = ws.cell(row=row, column=4, value=preco)
        c_preco.number_format = '#,##0.00'
        c_preco.border = thin_border
        ws.cell(row=row, column=5, value=qty).border = thin_border
        c_sub = ws.cell(row=row, column=6, value=subtotal)
        c_sub.number_format = '#,##0.00'
        c_sub.border = thin_border

    # Ajustar largura das colunas
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 8
    ws.column_dimensions["F"].width = 14

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def exportar_excel_todos(limit: int = 500) -> BytesIO:
    conn = _conn()
    rows = conn.execute(
        "SELECT * FROM recomendacoes ORDER BY id DESC LIMIT ?", (limit,)
    ).fetchall()
    conn.close()

    wb = openpyxl.Workbook()

    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="E63946", end_color="E63946", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Aba 1 - Resumo
    ws1 = wb.active
    ws1.title = "Resumo"
    headers1 = ["#", "Data", "Segmento", "Regiao", "UF", "Porte", "Total", "Qtd Produtos"]
    for col, h in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    for i, r in enumerate(rows, 2):
        d = dict(r)
        ws1.cell(row=i, column=1, value=d["id"]).border = thin_border
        ws1.cell(row=i, column=2, value=d["data"]).border = thin_border
        ws1.cell(row=i, column=3, value=d["segmento"]).border = thin_border
        ws1.cell(row=i, column=4, value=d["regiao"]).border = thin_border
        ws1.cell(row=i, column=5, value=d["uf"]).border = thin_border
        ws1.cell(row=i, column=6, value=d["porte"]).border = thin_border
        c_total = ws1.cell(row=i, column=7, value=d["total"])
        c_total.number_format = '#,##0.00'
        c_total.border = thin_border
        ws1.cell(row=i, column=8, value=d["qtd_produtos"]).border = thin_border

    ws1.column_dimensions["A"].width = 5
    ws1.column_dimensions["B"].width = 18
    ws1.column_dimensions["C"].width = 25
    ws1.column_dimensions["D"].width = 15
    ws1.column_dimensions["E"].width = 6
    ws1.column_dimensions["F"].width = 12
    ws1.column_dimensions["G"].width = 14
    ws1.column_dimensions["H"].width = 14

    # Aba 2 - Detalhes
    ws2 = wb.create_sheet("Detalhes")
    headers2 = ["Rec #", "Data", "Segmento", "Produto", "Formato", "Preco Unit.", "Qtd", "Subtotal"]
    for col, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    row_num = 2
    for r in rows:
        d = dict(r)
        produtos = json.loads(d.get("produtos_json", "[]"))
        for p in produtos:
            try:
                preco = float(p.get("preco", 0))
            except (ValueError, TypeError):
                preco = 0
            qty = int(p.get("quantidade", 1))
            ws2.cell(row=row_num, column=1, value=d["id"]).border = thin_border
            ws2.cell(row=row_num, column=2, value=d["data"]).border = thin_border
            ws2.cell(row=row_num, column=3, value=d["segmento"]).border = thin_border
            ws2.cell(row=row_num, column=4, value=p.get("nome", "")).border = thin_border
            ws2.cell(row=row_num, column=5, value=p.get("formato", "")).border = thin_border
            c_p = ws2.cell(row=row_num, column=6, value=preco)
            c_p.number_format = '#,##0.00'
            c_p.border = thin_border
            ws2.cell(row=row_num, column=7, value=qty).border = thin_border
            c_s = ws2.cell(row=row_num, column=8, value=preco * qty)
            c_s.number_format = '#,##0.00'
            c_s.border = thin_border
            row_num += 1

    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 18
    ws2.column_dimensions["C"].width = 25
    ws2.column_dimensions["D"].width = 40
    ws2.column_dimensions["E"].width = 15
    ws2.column_dimensions["F"].width = 14
    ws2.column_dimensions["G"].width = 8
    ws2.column_dimensions["H"].width = 14

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
